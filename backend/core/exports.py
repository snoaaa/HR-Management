"""
BN-19 / BN-126: export any list of the system as CSV, Excel or PDF.
"""

import csv
import io
import re
from decimal import Decimal

from django.http import HttpResponse


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    return value


def export_rows(headers, rows, filename, fmt="csv", title=None, subtitle=""):
    """
    headers: list of column titles.
    rows: iterable of lists (same order as headers).
    fmt: "csv", "xlsx" or "pdf".
    """
    rows = [[_cell(v) for v in row] for row in rows]
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = Workbook()
        ws = wb.active
        # Excel forbids some characters in sheet names.
        ws.title = re.sub(r"[\[\]:*?/\\]", "-", title or "Export")[:31]
        ws.append(headers)
        header_fill = PatternFill("solid", fgColor="714B67")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        zebra = PatternFill("solid", fgColor="F5F0F3")
        for index, row in enumerate(rows):
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.##"
                if index % 2:
                    cell.fill = zebra
        ws.freeze_panes = "A2"
        for column in ws.columns:
            width = max(len(str(c.value or "")) for c in column)
            ws.column_dimensions[column[0].column_letter].width = min(width + 2, 50)
        buffer = io.BytesIO()
        wb.save(buffer)
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
        return response
    if fmt == "pdf":
        from .pdf import render_table_pdf

        pdf = render_table_pdf(title or filename, headers, rows, subtitle)
        response = HttpResponse(pdf, content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
        return response

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(rows)
    # BOM so that Excel opens accented characters correctly.
    response = HttpResponse("\ufeff" + buffer.getvalue(), content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
    return response
